"""
CRM Bulk Import Service
Handles CSV and Excel parsing, data validation, preview, and transactional bulk importing.
"""

import io
import re
import csv
import openpyxl
from database.database import execute_query, get_db_cursor
from backend.services.activity_service import log_activity
from backend.services.audit_service import log_audit

MODULE_CONFIGS = {
    "customers": {
        "required": ["first_name", "last_name"],
        "columns": ["first_name", "last_name", "company_name", "email", "phone", "city", "state", "country", "industry", "status"],
        "sample": [
            ["Rajesh", "Sharma", "Sharma Logistics", "rajesh@sharmalogistics.in", "+91 98200 11223", "Ahmedabad", "Gujarat", "India", "Logistics", "active"],
            ["Priya", "Patel", "TechSolutions Ltd", "priya@techsolutions.com", "+91 98250 44556", "Surat", "Gujarat", "India", "Technology", "prospect"]
        ]
    },
    "leads": {
        "required": ["first_name", "last_name"],
        "columns": ["first_name", "last_name", "company_name", "email", "phone", "source", "status"],
        "sample": [
            ["Amit", "Verma", "Verma Enterprises", "amit@verma.in", "+91 98980 12345", "website", "new"],
            ["Neha", "Deshmukh", "Deshmukh Pharma", "neha@deshmukhpharma.com", "+91 98190 67890", "referral", "contacted"]
        ]
    },
    "products": {
        "required": ["name", "price"],
        "columns": ["product_code", "name", "category", "price", "tax_percentage", "discount_percentage", "stock", "status"],
        "sample": [
            ["PRD-101", "Enterprise CRM License", "Software", 45000.0, 18.0, 5.0, 100, "active"],
            ["PRD-102", "Annual Cloud Support Plan", "Services", 12000.0, 18.0, 0.0, 50, "active"]
        ]
    }
}

EMAIL_REGEX = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


def parse_import_file(file_obj, filename: str) -> tuple[list, list, str | None]:
    """Parse rows and headers from uploaded CSV or Excel file."""
    if not filename or "." not in filename:
        return [], [], "Invalid filename"

    ext = filename.rsplit(".", 1)[1].lower()
    headers = []
    rows = []

    try:
        if ext == "csv":
            content = file_obj.read().decode("utf-8-sig", errors="replace")
            reader = csv.reader(io.StringIO(content))
            for idx, r in enumerate(reader):
                clean_row = [str(c).strip() for c in r]
                if not any(clean_row):
                    continue
                if idx == 0:
                    headers = [c.lower().replace(" ", "_") for c in clean_row]
                else:
                    rows.append(clean_row)

        elif ext in ("xlsx", "xls"):
            wb = openpyxl.load_workbook(file_obj, data_only=True)
            ws = wb.active
            for idx, r in enumerate(ws.iter_rows(values_only=True)):
                clean_row = [str(c).strip() if c is not None else "" for c in r]
                if not any(clean_row):
                    continue
                if idx == 0:
                    headers = [c.lower().replace(" ", "_") for c in clean_row]
                else:
                    rows.append(clean_row)
        else:
            return [], [], f"Unsupported file format '.{ext}'. Upload a .csv or .xlsx file."

    except Exception as ex:
        return [], [], f"Failed to parse file: {str(ex)}"

    return headers, rows, None


def preview_import(file_obj, filename: str, module: str) -> tuple[dict | None, str | None]:
    """
    Validate columns and rows, providing a preview and error diagnostics before DB insertion.
    """
    mod = (module or "").lower().strip()
    if mod not in MODULE_CONFIGS:
        return None, f"Module '{module}' not supported for bulk import. Supported: customers, leads, products."

    cfg = MODULE_CONFIGS[mod]
    headers, rows, err = parse_import_file(file_obj, filename)
    if err:
        return None, err

    if not rows:
        return None, "The uploaded file contains no data rows"

    # Map headers to expected columns
    header_map = {}
    for idx, h in enumerate(headers):
        if h in cfg["columns"]:
            header_map[h] = idx

    # Check required headers
    missing_req = [req for req in cfg["required"] if req not in header_map]
    if missing_req:
        return None, f"Missing required column headers: {', '.join(missing_req)}. Expected headers include: {', '.join(cfg['columns'])}"

    validated_rows = []
    errors = []

    for row_num, r in enumerate(rows, start=2):
        row_dict = {}
        for col, idx in header_map.items():
            row_dict[col] = r[idx] if idx < len(r) else ""

        row_errors = []

        # Validate required fields
        for req in cfg["required"]:
            if not row_dict.get(req):
                row_errors.append(f"Missing {req.replace('_', ' ')}")

        # Validate email if provided
        email = row_dict.get("email")
        if email and not EMAIL_REGEX.match(email):
            row_errors.append(f"Invalid email '{email}'")

        # Validate numbers for products
        if mod == "products":
            try:
                price = float(row_dict.get("price", 0))
                if price < 0:
                    row_errors.append("Price cannot be negative")
                row_dict["price"] = price
            except ValueError:
                row_errors.append(f"Invalid price value '{row_dict.get('price')}'")

        is_valid = len(row_errors) == 0
        if not is_valid:
            errors.append({
                "row": row_num,
                "data": row_dict,
                "issues": row_errors
            })

        validated_rows.append({
            "row_number": row_num,
            "data": row_dict,
            "is_valid": is_valid,
            "errors": row_errors
        })

    valid_count = sum(1 for r in validated_rows if r["is_valid"])

    return {
        "module": mod,
        "filename": filename,
        "total_rows": len(rows),
        "valid_rows_count": valid_count,
        "invalid_rows_count": len(errors),
        "columns_mapped": list(header_map.keys()),
        "preview_rows": validated_rows[:20],
        "valid_rows": [r["data"] for r in validated_rows if r["is_valid"]],
        "errors": errors[:50]
    }, None


def confirm_import(module: str, rows_data: list, user_id: int | None = None, filename: str = "import_data") -> tuple[dict | None, str | None]:
    """
    Execute transactional insertion of validated rows.
    """
    mod = (module or "").lower().strip()
    if mod not in MODULE_CONFIGS:
        return None, f"Module '{module}' not supported"

    if not rows_data or not isinstance(rows_data, list):
        return None, "No data rows provided to import"

    imported_count = 0
    failed_count = 0
    error_notes = []

    try:
        with get_db_cursor(commit=True) as cursor:
            for item in rows_data:
                d = item.get("data", item)
                try:
                    if mod == "customers":
                        # Generate code if missing
                        code = d.get("customer_code") or f"CUST-{uuid.uuid4().hex[:6].upper()}"
                        status = d.get("status") or "prospect"
                        if status not in ("active", "inactive", "prospect", "blocked"):
                            status = "prospect"

                        sql = """
                            INSERT INTO customers (customer_code, first_name, last_name, company_name, email, phone, city, state, country, industry, status, created_at)
                            VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, NOW())
                        """
                        cursor.execute(sql, (
                            code, d.get("first_name"), d.get("last_name"), d.get("company_name") or None,
                            d.get("email") or None, d.get("phone") or None, d.get("city") or None,
                            d.get("state") or None, d.get("country") or None, d.get("industry") or None, status
                        ))
                        imported_count += 1

                    elif mod == "leads":
                        code = d.get("lead_code") or f"LEAD-{uuid.uuid4().hex[:6].upper()}"
                        status = d.get("status") or "new"
                        if status not in ("new", "contacted", "qualified", "proposal", "converted", "lost"):
                            status = "new"

                        sql = """
                            INSERT INTO leads (lead_code, first_name, last_name, company_name, email, phone, source, status, created_at)
                            VALUES (%s, %s, %s, %s, %s, %s, %s, %s, NOW())
                        """
                        cursor.execute(sql, (
                            code, d.get("first_name"), d.get("last_name"), d.get("company_name") or None,
                            d.get("email") or None, d.get("phone") or None, d.get("source") or "other", status
                        ))
                        imported_count += 1

                    elif mod == "products":
                        code = d.get("product_code") or f"PRD-{uuid.uuid4().hex[:6].upper()}"
                        status = d.get("status") or "active"
                        tax = float(d.get("tax_percentage") or 0.0)
                        disc = float(d.get("discount_percentage") or 0.0)
                        stock = int(d.get("stock") or 0)
                        price = float(d.get("price") or 0.0)

                        sql = """
                            INSERT INTO products (product_code, name, category, price, tax_percentage, discount_percentage, stock, status, created_at)
                            VALUES (%s, %s, %s, %s, %s, %s, %s, %s, NOW())
                        """
                        cursor.execute(sql, (
                            code, d.get("name"), d.get("category") or "General", price, tax, disc, stock, status
                        ))
                        imported_count += 1

                except Exception as row_ex:
                    failed_count += 1
                    error_notes.append(str(row_ex))

        # Log into import_jobs table
        sql_job = """
            INSERT INTO import_jobs (user_id, module_name, file_name, total_rows, imported_count, failed_count, status, error_summary, created_at)
            VALUES (%s, %s, %s, %s, %s, %s, %s, %s, NOW())
        """
        job_status = "completed" if failed_count == 0 else ("partially_completed" if imported_count > 0 else "failed")
        summary_text = f"Imported {imported_count}, Failed {failed_count}"
        if error_notes:
            summary_text += f" | First error: {error_notes[0][:100]}"

        execute_query(sql_job, (
            user_id, mod, filename, len(rows_data), imported_count, failed_count, job_status, summary_text
        ), commit=True)

        log_activity(user_id, mod, None, "bulk_import", f"Bulk imported {imported_count} {mod} records ({failed_count} failed)")
        log_audit(user_id, "bulk_import", mod, None, None, {"imported": imported_count, "failed": failed_count})

        return {
            "module": mod,
            "total": len(rows_data),
            "imported_count": imported_count,
            "failed_count": failed_count,
            "status": job_status
        }, None

    except Exception as ex:
        return None, f"Database transaction error during import: {str(ex)}"


def generate_sample_template(module: str, export_format: str = "csv") -> tuple[io.BytesIO | None, str, str, str | None]:
    """Generate downloadable template with column headers and sample rows."""
    mod = (module or "").lower().strip()
    if mod not in MODULE_CONFIGS:
        return None, "", "", f"Module '{module}' not found"

    cfg = MODULE_CONFIGS[mod]
    headers = cfg["columns"]
    samples = cfg["sample"]
    filename = f"sample_{mod}_import.{export_format}"

    if export_format == "xlsx":
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Import Template"
        ws.append(headers)
        for s in samples:
            ws.append(s)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf, filename, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", None
    else:
        buf = io.BytesIO()
        buf.write(b'\xef\xbb\xbf')
        text_buf = io.StringIO()
        writer = csv.writer(text_buf)
        writer.writerow(headers)
        for s in samples:
            writer.writerow(s)
        buf.write(text_buf.getvalue().encode('utf-8'))
        buf.seek(0)
        return buf, filename, "text/csv; charset=utf-8", None
