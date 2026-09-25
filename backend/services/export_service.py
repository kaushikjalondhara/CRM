"""
CRM Data Export Service
Exports CRM records to CSV, Excel (.xlsx), and PDF with formatting and filters.
"""

import io
import csv
import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from database.database import execute_query
from backend.services.pdf_service import generate_invoice_pdf


def export_data(module: str, export_format: str, filters: dict, current_user: dict) -> tuple[io.BytesIO | None, str, str, str | None]:
    """
    Main export dispatcher.
    Returns: (buffer, filename, mimetype, error)
    """
    mod = (module or "").lower().strip()
    fmt = (export_format or "csv").lower().strip()

    if fmt not in ("csv", "xlsx", "pdf"):
        return None, "", "", f"Unsupported format '{export_format}'. Must be csv, xlsx, or pdf."

    if mod == "invoices" and fmt == "pdf":
        inv_id = filters.get("invoice_id")
        if inv_id:
            buf, fname, err = generate_invoice_pdf(int(inv_id))
            return buf, fname, "application/pdf", err

    # Query module dataset
    headers, rows, filename_prefix, err = _fetch_module_dataset(mod, filters, current_user)
    if err:
        return None, "", "", err

    timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"{filename_prefix}_{timestamp}.{fmt}"

    if fmt == "csv":
        buf = _generate_csv(headers, rows)
        return buf, filename, "text/csv; charset=utf-8", None

    elif fmt == "xlsx":
        buf = _generate_excel(headers, rows, sheet_title=mod.capitalize())
        return buf, filename, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", None

    elif fmt == "pdf":
        buf = _generate_pdf_report(headers, rows, title=f"Apex CRM — {mod.capitalize()} Export")
        return buf, filename, "application/pdf", None

    return None, "", "", "Invalid format request"


def _fetch_module_dataset(module: str, filters: dict, current_user: dict) -> tuple[list, list, str, str | None]:
    """
    Fetch filtered dataset based on module name.
    """
    search = filters.get("search", "").strip()
    status = filters.get("status")
    start_date = filters.get("start_date")
    end_date = filters.get("end_date")
    assigned_to = filters.get("assigned_to")

    where = []
    params = []

    if module == "customers":
        if search:
            where.append("(c.customer_code LIKE %s OR c.first_name LIKE %s OR c.last_name LIKE %s OR c.company_name LIKE %s OR c.email LIKE %s OR c.phone LIKE %s)")
            t = f"%{search}%"
            params.extend([t, t, t, t, t, t])
        if status:
            where.append("c.status = %s")
            params.append(status)
        if assigned_to:
            where.append("c.assigned_to = %s")
            params.append(assigned_to)

        where_sql = f"WHERE {' AND '.join(where)}" if where else ""
        sql = f"""
            SELECT c.customer_code, CONCAT(c.first_name, ' ', c.last_name) AS full_name,
                   c.company_name, c.email, c.phone, c.city, c.state, c.status,
                   CONCAT(u.first_name, ' ', u.last_name) AS assigned_to_name,
                   DATE_FORMAT(c.created_at, '%Y-%m-%d') AS created_date
            FROM customers c
            LEFT JOIN users u ON c.assigned_to = u.id
            {where_sql}
            ORDER BY c.id DESC LIMIT 5000
        """
        rows, err = execute_query(sql, tuple(params), fetch_all=True)
        headers = ["Customer Code", "Contact Name", "Company", "Email", "Phone", "City", "State", "Status", "Assigned To", "Created Date"]
        data = [[r.get(k, '') or '' for k in ['customer_code', 'full_name', 'company_name', 'email', 'phone', 'city', 'state', 'status', 'assigned_to_name', 'created_date']] for r in (rows or [])]
        return headers, data, "Customers_Export", err

    elif module == "leads":
        if search:
            where.append("(l.lead_code LIKE %s OR l.first_name LIKE %s OR l.last_name LIKE %s OR l.company_name LIKE %s OR l.email LIKE %s)")
            t = f"%{search}%"
            params.extend([t, t, t, t, t])
        if status:
            where.append("l.status = %s")
            params.append(status)

        where_sql = f"WHERE {' AND '.join(where)}" if where else ""
        sql = f"""
            SELECT l.lead_code, CONCAT(l.first_name, ' ', l.last_name) AS full_name,
                   l.company_name, l.email, l.phone, l.source, l.status,
                   CONCAT(u.first_name, ' ', u.last_name) AS assigned_to_name,
                   DATE_FORMAT(l.created_at, '%Y-%m-%d') AS created_date
            FROM leads l
            LEFT JOIN users u ON l.assigned_to = u.id
            {where_sql}
            ORDER BY l.id DESC LIMIT 5000
        """
        rows, err = execute_query(sql, tuple(params), fetch_all=True)
        headers = ["Lead Code", "Name", "Company", "Email", "Phone", "Source", "Status", "Assigned To", "Created Date"]
        data = [[r.get(k, '') or '' for k in ['lead_code', 'full_name', 'company_name', 'email', 'phone', 'source', 'status', 'assigned_to_name', 'created_date']] for r in (rows or [])]
        return headers, data, "Leads_Export", err

    elif module == "payments":
        if search:
            where.append("(p.transaction_reference LIKE %s OR i.invoice_number LIKE %s)")
            t = f"%{search}%"
            params.extend([t, t])

        where_sql = f"WHERE {' AND '.join(where)}" if where else ""
        sql = f"""
            SELECT p.transaction_reference, i.invoice_number,
                   CONCAT(c.first_name, ' ', c.last_name) AS customer_name,
                   p.amount, p.payment_method, DATE_FORMAT(p.payment_date, '%Y-%m-%d') AS p_date,
                   p.notes
            FROM payments p
            JOIN invoices i ON p.invoice_id = i.id
            JOIN customers c ON p.customer_id = c.id
            {where_sql}
            ORDER BY p.id DESC LIMIT 5000
        """
        rows, err = execute_query(sql, tuple(params), fetch_all=True)
        headers = ["Transaction Ref", "Invoice #", "Customer", "Amount (₹)", "Payment Method", "Date", "Notes"]
        data = [[r.get(k, '') or '' for k in ['transaction_reference', 'invoice_number', 'customer_name', 'amount', 'payment_method', 'p_date', 'notes']] for r in (rows or [])]
        return headers, data, "Payments_Export", err

    elif module == "products":
        if search:
            where.append("(name LIKE %s OR product_code LIKE %s OR category LIKE %s)")
            t = f"%{search}%"
            params.extend([t, t, t])

        where_sql = f"WHERE {' AND '.join(where)}" if where else ""
        sql = f"""
            SELECT product_code, name, category, price, tax_percentage, discount_percentage, stock, status
            FROM products
            {where_sql}
            ORDER BY id DESC LIMIT 5000
        """
        rows, err = execute_query(sql, tuple(params), fetch_all=True)
        headers = ["Product Code", "Name", "Category", "Price (₹)", "Tax %", "Discount %", "Stock", "Status"]
        data = [[r.get(k, '') or '' for k in ['product_code', 'name', 'category', 'price', 'tax_percentage', 'discount_percentage', 'stock', 'status']] for r in (rows or [])]
        return headers, data, "Products_Catalog_Export", err

    elif module == "invoices":
        if search:
            where.append("(i.invoice_number LIKE %s OR c.first_name LIKE %s OR c.last_name LIKE %s)")
            t = f"%{search}%"
            params.extend([t, t, t])
        if status:
            where.append("i.status = %s")
            params.append(status)

        where_sql = f"WHERE {' AND '.join(where)}" if where else ""
        sql = f"""
            SELECT i.invoice_number, CONCAT(c.first_name, ' ', c.last_name) AS customer_name,
                   DATE_FORMAT(i.invoice_date, '%Y-%m-%d') AS inv_date,
                   DATE_FORMAT(i.due_date, '%Y-%m-%d') AS due_date,
                   i.subtotal, i.tax_amount, i.discount_amount, i.total_amount,
                   i.paid_amount, i.remaining_amount, i.status
            FROM invoices i
            JOIN customers c ON i.customer_id = c.id
            {where_sql}
            ORDER BY i.id DESC LIMIT 5000
        """
        rows, err = execute_query(sql, tuple(params), fetch_all=True)
        headers = ["Invoice #", "Customer", "Date", "Due Date", "Subtotal", "Tax", "Discount", "Total (₹)", "Paid (₹)", "Balance (₹)", "Status"]
        data = [[r.get(k, '') or '' for k in ['invoice_number', 'customer_name', 'inv_date', 'due_date', 'subtotal', 'tax_amount', 'discount_amount', 'total_amount', 'paid_amount', 'remaining_amount', 'status']] for r in (rows or [])]
        return headers, data, "Invoices_Export", err

    elif module == "reports":
        # General Executive Summary
        sql = """
            SELECT 'Customers' as Metric, COUNT(*) as Count, 'Registered Clients' as Note FROM customers
            UNION ALL
            SELECT 'Leads', COUNT(*), 'Pipeline Leads' FROM leads
            UNION ALL
            SELECT 'Deals', COUNT(*), 'Active Opportunities' FROM deals
            UNION ALL
            SELECT 'Invoices', COUNT(*), CONCAT('Total: ₹', ROUND(COALESCE(SUM(total_amount), 0), 2)) FROM invoices
            UNION ALL
            SELECT 'Payments', COUNT(*), CONCAT('Collected: ₹', ROUND(COALESCE(SUM(amount), 0), 2)) FROM payments
        """
        rows, err = execute_query(sql, fetch_all=True)
        headers = ["Metric Category", "Total Count", "Financial Note"]
        data = [[r.get(k, '') for k in ['Metric', 'Count', 'Note']] for r in (rows or [])]
        return headers, data, "CRM_Executive_Report", err

    return [], [], "Export", f"Module '{module}' not recognized for export"


def _generate_csv(headers: list, rows: list) -> io.BytesIO:
    """Generate UTF-8 CSV buffer with BOM for Excel compatibility."""
    buffer = io.BytesIO()
    # Write UTF-8 BOM
    buffer.write(b'\xef\xbb\xbf')
    text_buffer = io.StringIO()
    writer = csv.writer(text_buffer)
    writer.writerow(headers)
    for row in rows:
        writer.writerow(row)
    buffer.write(text_buffer.getvalue().encode('utf-8'))
    buffer.seek(0)
    return buffer


def _generate_excel(headers: list, rows: list, sheet_title: str = "Export") -> io.BytesIO:
    """Generate professional styled Excel file using openpyxl."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]

    # Header styling: Primary Blue #2563eb, Bold White text
    header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    ws.append(headers)
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.row_dimensions[1].height = 24

    for row_idx, row_data in enumerate(rows, start=2):
        ws.append(row_data)
        ws.row_dimensions[row_idx].height = 20
        for col_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 40)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def _generate_pdf_report(headers: list, rows: list, title: str = "CRM Report") -> io.BytesIO:
    """Generate printable PDF table using ReportLab."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
    styles = getSampleStyleSheet()

    story = [
        Paragraph(f"<b>{title}</b>", ParagraphStyle('ReportTitle', parent=styles['Heading1'], fontSize=16, textColor=colors.HexColor('#1e3a8a'))),
        Paragraph(f"Exported on: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", ParagraphStyle('ReportDate', parent=styles['Normal'], fontSize=9, textColor=colors.HexColor('#64748b'))),
        Spacer(1, 14),
        HRFlowable(width="100%", thickness=1, color=colors.HexColor("#cbd5e1"), spaceAfter=14)
    ]

    table_data = [[Paragraph(f"<b>{h}</b>", styles['Normal']) for h in headers]]
    for r in rows[:100]:  # Limit 100 rows in PDF summary to prevent memory overflow
        table_data.append([Paragraph(str(c), styles['Normal']) for c in r])

    t = Table(table_data)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#f1f5f9')),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    story.append(t)

    doc.build(story)
    buffer.seek(0)
    return buffer
